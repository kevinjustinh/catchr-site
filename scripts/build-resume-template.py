"""Builds downloads/catchr-resume-template.docx.

Shaped for applicant tracking systems and for Catchr's Tailor feature (see the
extension's background.js tailor prompt and applyBulletReorders):
- one column, no tables, text boxes or images;
- each job is one unique heading line "Title, Company<TAB>dates", so Tailor can
  find that job's bullets and reorder them (its matcher folds the tab to a space);
- real one-level bullets; skills as "Category: item, item" (Tailor edits only
  after the colon); the summary is a single paragraph.

Run: python3 scripts/build-resume-template.py  (needs python-docx)
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "downloads" / "catchr-resume-template.docx"
FONT = "Lato"
INK = RGBColor(0x2C, 0x2C, 0x2A)
STONE = RGBColor(0x5E, 0x56, 0x4C)  # darker than the site's stone: stays legible printed or scanned
FOREST = RGBColor(0x2C, 0x5F, 0x4A)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.15
bullet_style = doc.styles["List Bullet"]
bullet_style.font.name = FONT
bullet_style.paragraph_format.space_after = Pt(2)


def para(text="", size=None, bold=False, color=None, after=0, before=0, caps=False, spacing=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if caps:
            r.font.all_caps = True
        if spacing:
            rpr = r._element.get_or_add_rPr()
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(spacing))
            rpr.append(sp)
    return p


def rule_below(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "2"), ("color", "B8C9BF")):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    ppr.append(bdr)


def heading(text):
    # A real Heading 2 (outline, screen readers), restyled to match the page.
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold, r.font.all_caps, r.font.size = True, True, Pt(11)
    r.font.name, r.font.color.rgb = FONT, FOREST
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), "20")
    r._element.get_or_add_rPr().append(sp)
    rule_below(p)


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)


def job(title, dates, where):
    p = para(before=8)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(title).bold = True
    p.add_run("\t" + dates)
    para(where, color=STONE, size=10, after=3)


# Everything below is fictional: Contoso and Fabrikam are Microsoft's stock
# example companies, 555-01xx numbers are reserved for fiction, and
# example.com is reserved for documentation.
para("Jane Doe", size=22, bold=True, after=2)
para("Austin, TX · jane.doe@example.com · (512) 555-0147 · linkedin.com/in/your-profile",
     size=10, color=STONE, after=4)

heading("Summary")
para("Product manager with 7 years in B2B SaaS, most recently owning self-serve growth at an analytics "
     "company. Pairs customer research with product data to decide what to build, and what not to. "
     "Looking for a senior product role on a growth or platform team.")

heading("Experience")
job("Senior Product Manager, Contoso Analytics", "Mar 2022 to Present", "Austin, TX · Hybrid")
bullet("Led a redesign of self-serve onboarding with 1 designer and 6 engineers, raising trial-to-paid conversion from 9% to 13% in two quarters.")
bullet("Cut the roadmap from 14 initiatives to 5 after 60 customer interviews traced churn to reporting gaps, reducing customer churn 18% year over year.")
bullet("Introduced feature flags and staged rollouts with engineering, shortening the release cycle from 6 weeks to 2.")
bullet("Defined the product's activation metric with data science; it became the company-wide goal used in quarterly planning.")

job("Product Manager, Fabrikam Field Services", "Jun 2019 to Feb 2022", "Denver, CO")
bullet("Launched an offline-first mobile app for 30,000 field technicians, reaching a 4.7 App Store rating within six months.")
bullet("Partnered with sales on packaging and pricing research that raised average contract value 12%.")
bullet("Wrote specs and acceptance criteria for 6 engineers and 1 designer, shipping 3 major releases a year.")

job("Associate Product Manager, Fabrikam Field Services", "Jul 2017 to May 2019", "Denver, CO")
bullet("Ran 40 A/B tests on the signup flow and shipped the 5 winners, adding about 2,100 signups a month.")
bullet("Built the team's first SQL funnel dashboard, replacing a manual weekly report.")

heading("Skills")
for cat, items in (("Product", "Roadmapping, customer discovery, A/B testing, pricing and packaging, PRDs"),
                   ("Tools", "Jira, Figma, Amplitude, Mixpanel, SQL, Looker"),
                   ("Leadership", "Cross-functional leadership, stakeholder communication, mentoring")):
    p = para(after=2)
    r = p.add_run(f"{cat}: ")
    r.bold = True
    p.add_run(items)

heading("Education")
p = para()
p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
p.add_run("B.S. Business Administration, State University").bold = True
p.add_run("\tMay 2017")

doc.core_properties.author = "Catchr"
doc.core_properties.title = "Resume template"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT}")
